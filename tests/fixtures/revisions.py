"""Word files with tracked changes, for the tests of accept_all and reject_all.

``write_revisions`` writes one of each kind of revision Word stores, every one in a sentence that
says which side it belongs to:

- inserted and deleted text: "The Buyer shall pay EUR [1,000 -> 2,000] within [30 -> 45] days.";
- a formatting change that made the word "bold" bold (its old properties are plain), a paragraph
  change that centred a right-aligned paragraph (its mark is italic), and a numbering change;
- a sentence moved from clause A to clause B, with the move's range markers;
- an inserted paragraph mark ("First half, " / "second half.") and a deleted one ("Joined " /
  "together.", the second paragraph justified);
- a text box whose two copies, DrawingML in mc:Choice and VML in mc:Fallback, hold the same
  change "Box [old -> new]";
- deleted text outside any deletion, as a damaged file might hold it: "Stray text is [not] shown.";
- a table whose properties changed, with a header row whose row properties and property
  exceptions changed, an inserted and a deleted row, an inserted and a deleted cell, a merged
  cell and a cell whose shading changed;
- a change of page size in the section, and changes in the header and the footer.

``write_paragraph_marks`` writes three paragraphs whose middle paragraph mark is a tracked
insertion or deletion. python-docx makes the package (body, header, footer, styles); the
revisions, which it has no API for, are written as WordprocessingML with lxml, the way Word writes
them.
"""

from __future__ import annotations

import itertools
from pathlib import Path
from typing import Literal

import pytest
from lxml import etree

from tests.fixtures.docx import XMLNS
from themask.docx.package import DocxPackage

MarkKind = Literal["ins", "del"]

_REVIEWER = 'w:author="Reviewer" w:date="2026-10-01T09:00:00Z"'
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class _Marks:
    """Writes revision marks, each with an id of its own, as Word gives them."""

    def __init__(self) -> None:
        self._ids = itertools.count(1)

    def attrs(self) -> str:
        """The attributes of one revision: its id, the reviewer and the date."""
        return f'w:id="{next(self._ids)}" {_REVIEWER}'

    def inserted(self, text: str) -> str:
        return f"<w:ins {self.attrs()}>{run(text)}</w:ins>"

    def deleted(self, text: str) -> str:
        return f"<w:del {self.attrs()}>{deleted_run(text)}</w:del>"

    def mark(self, kind: MarkKind) -> str:
        """The run properties of a paragraph mark inserted or deleted as a tracked change."""
        return f"<w:rPr><w:{kind} {self.attrs()}/></w:rPr>"


def run(text: str, props: str = "") -> str:
    return f'<w:r>{props}<w:t xml:space="preserve">{text}</w:t></w:r>'


def deleted_run(text: str) -> str:
    return f'<w:r><w:delText xml:space="preserve">{text}</w:delText></w:r>'


def para(content: str, props: str = "") -> str:
    return f"<w:p>{f'<w:pPr>{props}</w:pPr>' if props else ''}{content}</w:p>"


def _element(xml: str) -> etree._Element:
    return etree.fromstring(f"<w:root {XMLNS}>{xml}</w:root>")[0]


def _cell(content: str, props: str = "") -> str:
    return f'<w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/>{props}</w:tcPr>{content}</w:tc>'


def _text_box(m: _Marks) -> str:
    """A run holding a text box, its paragraph changed the same way in both copies."""
    inside = para(run("Box ") + m.deleted("old") + m.inserted("new"))
    return (
        '<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f"<wps:wsp><wps:txbx><w:txbxContent>{inside}</w:txbxContent></wps:txbx></wps:wsp>"
        "</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{inside}</w:txbxContent>"
        "</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    )


def _row(props: str, *cells: str) -> str:
    return f"<w:tr>{props}{''.join(cells)}</w:tr>"


def _table(m: _Marks) -> str:
    def tracked_cell(text: str, kind: MarkKind) -> str:
        # Word marks the content and the paragraph mark of every cell of a row it tracks.
        content = m.inserted(text) if kind == "ins" else m.deleted(text)
        return _cell(para(content, m.mark(kind)))

    rows = [
        _row(  # the header row: its properties and its table property exceptions changed
            "<w:tblPrEx><w:tblBorders>"
            '<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/></w:tblBorders>'
            f"<w:tblPrExChange {m.attrs()}><w:tblPrEx/></w:tblPrExChange></w:tblPrEx>"
            f"<w:trPr><w:tblHeader/><w:trPrChange {m.attrs()}><w:trPr/></w:trPrChange></w:trPr>",
            _cell(para(run("Item"))),
            _cell(para(run("Price"))),
        ),
        _row(
            f"<w:trPr><w:ins {m.attrs()}/></w:trPr>",
            tracked_cell("Added row", "ins"),
            tracked_cell("EUR 500", "ins"),
        ),
        _row(
            f"<w:trPr><w:del {m.attrs()}/></w:trPr>",
            tracked_cell("Removed row", "del"),
            tracked_cell("EUR 900", "del"),
        ),
        _row(
            "",
            _cell(para(run("Kept cell"))),
            _cell(para(run("Inserted cell")), f"<w:cellIns {m.attrs()}/>"),
            _cell(para(run("Deleted cell")), f"<w:cellDel {m.attrs()}/>"),
        ),
        _row(
            "",
            _cell(
                para(run("Merged cell")),
                f'<w:vMerge w:val="restart"/><w:cellMerge {m.attrs()} w:vMerge="rest"/>',
            ),
            _cell(
                para(run("Shaded cell")),
                '<w:shd w:val="clear" w:color="auto" w:fill="FFFF00"/>'
                f'<w:tcPrChange {m.attrs()}><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>'
                "</w:tcPrChange>",
            ),
        ),
    ]
    return (
        '<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="pct"/>'
        f'<w:tblPrChange {m.attrs()}><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        "</w:tblPrChange></w:tblPr>"
        '<w:tblGrid><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/>'
        f"</w:tblGrid>{''.join(rows)}</w:tbl>"
    )


def _body(m: _Marks) -> list[str]:
    """The blocks of the body, in order, each one holding one kind of revision or two."""
    move = f'w:name="move1" {_REVIEWER}'
    return [
        para(
            run("The Buyer shall pay EUR ")
            + m.deleted("1,000")
            + m.inserted("2,000")
            + run(" within ")
            + m.deleted("30")
            + m.inserted("45")
            + run(" days.")
        ),
        para(
            run("Made ")
            + run("bold", f"<w:rPr><w:b/><w:rPrChange {m.attrs()}><w:rPr/></w:rPrChange></w:rPr>")
            + run(" by a formatting change.")
        ),
        para(
            run("Centred by a paragraph change."),
            '<w:jc w:val="center"/><w:rPr><w:i/></w:rPr>'
            f'<w:pPrChange {m.attrs()}><w:pPr><w:jc w:val="right"/></w:pPr></w:pPrChange>',
        ),
        para(
            run("Clause A. ")
            + f'<w:moveFromRangeStart w:id="901" {move}/>'
            + f"<w:moveFrom {m.attrs()}>{run('Moved sentence. ')}</w:moveFrom>"
            + '<w:moveFromRangeEnd w:id="901"/>'
            + run("End of A.")
        ),
        para(
            run("Clause B. ")
            + f'<w:moveToRangeStart w:id="902" {move}/>'
            + f"<w:moveTo {m.attrs()}>{run('Moved sentence. ')}</w:moveTo>"
            + '<w:moveToRangeEnd w:id="902"/>'
            + run("End of B.")
        ),
        para(run("First half, "), m.mark("ins")),
        para(run("second half.")),
        para(run("Joined "), '<w:jc w:val="right"/>' + m.mark("del")),
        para(run("together."), '<w:jc w:val="both"/>'),
        para(
            run("A numbered item."),
            '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/>'
            f'<w:numberingChange {m.attrs()} w:original="%1."/></w:numPr>',
        ),
        para(run("Text box: ") + _text_box(m)),
        para(
            '<w:r><w:t xml:space="preserve">Stray text is </w:t>'
            '<w:delText xml:space="preserve">not </w:delText><w:t>shown.</w:t></w:r>'
        ),
        _table(m),
        para(run("End.")),
    ]


def write_revisions(path: Path) -> Path:
    """Writes the Word file with one of each kind of revision (see the module) at ``path``."""
    import docx  # here, not at the top: only the tests that ask for this file need python-docx

    document = docx.Document()
    m = _Marks()
    body = document.element.body
    section_props = body.find(f"{{{_W}}}sectPr")
    for block in _body(m):
        section_props.addprevious(_element(block))
    section_props.append(
        _element(
            f'<w:sectPrChange {m.attrs()}><w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
            '<w:pgMar w:top="1417" w:right="1417" w:bottom="1134" w:left="1417" w:header="708"'
            ' w:footer="708" w:gutter="0"/></w:sectPr></w:sectPrChange>'
        )
    )
    section = document.sections[0]
    for story, content in (
        (section.header, run("Draft ") + m.deleted("v1") + m.inserted("v2")),
        (section.footer, run("Confidential") + m.inserted(" draft")),
    ):
        old = story.paragraphs[0]._p
        old.getparent().replace(old, _element(para(content)))
    document.save(str(path))
    return path


def write_paragraph_marks(path: Path, kind: MarkKind) -> Path:
    """Writes "One", "Two", "Three" as three paragraphs, the mark of "Two" inserted or deleted."""
    import docx

    document = docx.Document()
    m = _Marks()
    section_props = document.element.body.find(f"{{{_W}}}sectPr")
    for block in (para(run("One")), para(run("Two"), m.mark(kind)), para(run("Three"))):
        section_props.addprevious(_element(block))
    document.save(str(path))
    return path


@pytest.fixture(scope="session")
def revisions_docx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """The Word file with one of each kind of revision, written once per test session."""
    return write_revisions(tmp_path_factory.mktemp("revisions") / "revisions.docx")


@pytest.fixture
def rev_pkg(revisions_docx: Path) -> DocxPackage:
    """The file with one of each kind of revision, opened afresh for each test."""
    return DocxPackage.open(revisions_docx)


@pytest.fixture
def rev_pkg_marks(tmp_path: Path) -> DocxPackage:
    """Three paragraphs, "One", "Two" and "Three", the mark of "Two" deleted as a tracked change.

    Accepting the deletion joins "Two" with "Three"; rejecting it keeps the three paragraphs.
    ``write_paragraph_marks(path, "ins")`` writes the same file with the mark inserted instead.
    """
    return DocxPackage.open(write_paragraph_marks(tmp_path / "marks.docx", "del"))
