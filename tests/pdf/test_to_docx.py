"""PDF to Word: a fresh Word file holding the text of a PDF, and nothing else of it.

The brief's test checks that every secret of supply-01 reaches the Word file. The others check how
the blocks are written (headings, paragraphs, a table with a bold first row), where a note stands
for a page that is only a picture, that nothing but text comes from the PDF and none of
python-docx's own properties stay, and that characters a Word file cannot hold are left out.
"""

from __future__ import annotations

import json
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

from tests.fixtures.pdf import Line, MakePdf
from themask.docx.package import DocxPackage
from themask.docx.reader import extract_text
from themask.pdf.extractor import extract_blocks
from themask.pdf.to_docx import IMAGE_PAGE_NOTE, pdf_to_docx
from themask.textnorm import normalize


def test_to_docx_keeps_every_truth_value(corpus: Path, tmp_path: Path) -> None:
    out = tmp_path / "s.docx"
    pdf_to_docx(corpus / "supply-01.pdf", out)
    text = normalize(extract_text(DocxPackage.open(out)).text)
    for t in json.loads((corpus / "supply-01.truth.json").read_text()):
        assert normalize(t["value"]) in text, t["value"]


def test_headings_paragraphs_and_a_table(corpus: Path, tmp_path: Path) -> None:
    out = tmp_path / "s.docx"
    report = pdf_to_docx(corpus / "supply-01.pdf", out)
    blocks, read = extract_blocks(corpus / "supply-01.pdf")
    assert report == read
    document = docx.Document(str(out))
    written = [(p.style.name if p.style else None, p.text) for p in document.paragraphs]
    expected_style = {"heading": "Heading 2", "paragraph": "Normal"}
    assert written == [
        ("Heading 1" if b.size >= 14 else expected_style[b.kind], b.text)
        for b in blocks
        if b.kind != "table"
    ]
    assert written[2] == ("Heading 1", "SUPPLY AGREEMENT")  # 16 pt; the clause headings are 12 pt
    (table,) = document.tables
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    assert rows == next(b.rows for b in blocks if b.kind == "table")
    for number, row in enumerate(table.rows):
        runs = [run for cell in row.cells for run in cell.paragraphs[0].runs]
        assert runs and all(run.bold is (True if number == 0 else None) for run in runs)


def test_a_note_stands_for_each_image_page(corpus: Path, tmp_path: Path) -> None:
    out = tmp_path / "m.docx"
    report = pdf_to_docx(corpus / "mixed-01.pdf", out)
    assert report.image_only_pages == [2]
    note = "[page 2 is an image \N{EM DASH} not included]"
    assert IMAGE_PAGE_NOTE.format(page=2) == note
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    assert texts.count(note) == 1
    blocks, _ = extract_blocks(corpus / "mixed-01.pdf")
    before = [b.text for b in blocks if b.page == 1 and b.kind != "table"][-1]
    after = next(b.text for b in blocks if b.page == 3 and b.kind != "table")
    at = texts.index(note)
    assert texts[at - 1 : at + 2] == [before, note, after]


def test_nothing_but_the_text_comes_from_the_pdf(corpus: Path, tmp_path: Path) -> None:
    """The PDF's title, author and the like stay behind, and python-docx's own are cleared."""
    out = tmp_path / "s.docx"
    pdf_to_docx(corpus / "supply-01.pdf", out)
    props = docx.Document(str(out)).core_properties
    texts = (props.author, props.comments, props.keywords, props.last_modified_by, props.subject)
    assert texts == ("",) * 5 and props.title == ""
    core = DocxPackage.open(out).raw("docProps/core.xml").decode()
    assert "python-docx" not in core and "Supply Agreement" not in core


def test_characters_word_cannot_hold_are_left_out(make_pdf: MakePdf, tmp_path: Path) -> None:
    pdf = make_pdf("nul.pdf", [[Line(100, "Alpha\x01Beta", embedded=True)]])
    out = tmp_path / "n.docx"
    pdf_to_docx(pdf, out)
    paragraphs: list[Paragraph] = docx.Document(str(out)).paragraphs
    assert [p.text for p in paragraphs] == ["AlphaBeta"]
